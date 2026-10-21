#!/usr/bin/env python3
"""Build conservative cross-store product matches from a catalog workbook.

The matcher deliberately separates candidate generation from publication.  It
normalizes Spanish retail wording, checks package measurements and rejects
common variant conflicts before creating one-product-per-store groups.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Iterable
from urllib.parse import unquote, urlparse


ACTIVE_STORES = ("TaTa", "Disco", "El Dorado")
SOURCE_SHEETS = ACTIVE_STORES

IN_SCOPE_SOURCE_CATEGORIES = {
    "TaTa": {"almacen", "bebidas", "frescos", "limpieza", "congelados"},
    "Disco": {"almacen", "bebidas", "frescos", "perfumeria y limpieza"},
    "El Dorado": {"almacen", "bebidas", "frescos", "limpieza e higiene", "kiosko"},
}

STOPWORDS = {
    "a", "al", "con", "de", "del", "el", "en", "la", "las", "los", "para",
    "por", "un", "una", "y", "x", "producto", "pack", "paquete", "pote",
    "frasco", "botella", "pet", "caja", "bolsa", "sachet", "unidad", "unidades",
}

GENERIC_TOKENS = {
    "aceite", "agua", "alimento", "bebida", "carne", "cafe", "detergente",
    "galleta", "jabon", "jugo", "leche", "limpiador", "papel", "queso",
    "refresco", "shampoo", "yogur",
}

FLAVORS = {
    "ananá", "banana", "chocolate", "coco", "durazno", "frutilla", "limon",
    "lavanda", "mandarina", "mango", "manzana", "menta", "naranja", "pomelo", "tropical",
    "vainilla", "dulcedeleche", "frutosrojos", "multifruta", "original",
}

WINE_VARIANTS = {
    "cabernet", "carmenere", "chardonnay", "malbec", "merlot", "moscatel",
    "pinot", "sauvignon", "tannat",
}

BUNDLE_EXTRAS = {"aparato", "copa", "gratis", "mayonnaise", "regalo", "yerbera"}

ONE_SIDED_CRITICAL_VARIANTS = {
    "zero", "light", "descremada", "semidescremada", "congas", "singas", "consal", "sinsal",
}

VARIANT_FAMILIES = (
    frozenset({"zero", "light", "original"}),
    frozenset({"entera", "integral", "descremada", "semidescremada"}),
    frozenset({"congas", "singas"}),
    frozenset({"consal", "sinsal"}),
    frozenset({"liquido", "polvo"}),
    frozenset({"recarga", "aerosol", "gatillo"}),
)

CLEANING_TERMS = {
    "absorbe", "ambientador", "aromatizante", "balde", "bolsa", "brillo", "cera",
    "cepillo", "desengrasante", "desinfectante", "detergente", "escoba", "esponja",
    "guante", "hipoclorito", "insecticida", "jabon", "lampazo", "lavandina",
    "lavavajilla", "limpiador", "limpieza", "lustramueble", "papel", "repelente",
    "servilleta", "suavizante", "toalla", "trapo",
}

DAIRY_TERMS = {
    "chocolatada", "colet", "crema", "flan", "helado", "lactolate", "leche",
    "manteca", "margarina", "muzzarella", "postre", "queso", "ricota", "yogur",
}

MEAT_TERMS = {
    "asado", "bife", "bondiola", "carne", "carre", "cerdo", "chorizo", "costilla",
    "entraña", "hamburguesa", "jamon", "matambre", "milanesa", "mortadela", "pollo",
    "pulpón", "salame", "salchicha", "vacío",
}

PREPARED_FOOD_TERMS = {
    "croqueta", "croquetas", "empanada", "empanadas", "pan", "pasta", "pastas",
    "pascualina", "pizza", "raviol", "ravioles", "sandwich", "tarta", "tartas",
}


@dataclass(frozen=True)
class Measurement:
    dimension: str
    value: float


@dataclass
class ProductRecord:
    record_id: int
    store: str
    source_category: str
    title: str
    url: str
    normalized: str
    tokens: list[str]
    measurements: list[Measurement]


@dataclass(frozen=True)
class MatchEdge:
    left_id: int
    right_id: int
    score: float
    quantity_state: str
    reason: str


def strip_accents(value: str) -> str:
    return "".join(
        char for char in unicodedata.normalize("NFKD", value)
        if not unicodedata.combining(char)
    )


def repair_text(value: object) -> str:
    text = str(value or "").strip()
    if "Ã" in text:
        try:
            text = text.encode("latin-1").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return text


def title_from_url(url: str) -> str:
    path = [part for part in urlparse(url).path.split("/") if part]
    if not path:
        return ""
    slug = path[-2] if path[-1].lower() in {"p", "producto"} and len(path) > 1 else path[-1]
    slug = re.sub(r"\.producto$", "", slug, flags=re.IGNORECASE)
    return unquote(slug).replace("-", " ")


def semantic_normalize(value: str) -> str:
    text = strip_accents(repair_text(value)).lower()
    replacements = (
        (r"\b(?:sin|cero|zero)\s+azucar(?:es)?\b", " zero "),
        (r"\b0\s*%\s*(?:azucar(?:es)?)?\b", " zero "),
        (r"\bcero\b", " zero "),
        (r"\bsin\s+gas\b", " singas "),
        (r"\bcon\s+gas\b", " congas "),
        (r"\bsin\s+sal\b", " sinsal "),
        (r"\bcon\s+sal\b", " consal "),
        (r"\bsin\s+lactosa\b|\bdeslactosad[oa]\b", " sinlactosa "),
        (r"\bsin\s+gluten\b", " singluten "),
        (r"\bdulce\s+de\s+leche\b", " dulcedeleche "),
        (r"\bfrutos?\s+rojos?\b", " frutosrojos "),
        (r"\b(?:ultra\s*pasteurizad[oa]|uht)\b", " ultra "),
        (r"\bsemidescremad[oa]\b", " semidescremada "),
        (r"\bdescremad[oa]\b", " descremada "),
        (r"\benter[oa]\b", " entera "),
        (r"\bmozzarella\b", " muzzarella "),
        (r"\bmayonesa\b", " mayonnaise "),
        (r"\bsabor\s+original\b", " original "),
        (r"\b(?:litros?|lts?|lt)\b", " l "),
        (r"\b(?:mililitros?|mls?)\b", " ml "),
        (r"\b(?:centimetros?\s*cubicos?|cc)\b", " ml "),
        (r"\b(?:kilogramos?|kilos?|kgs?)\b", " kg "),
        (r"\b(?:gramos?|grs?|gr)\b", " g "),
        (r"\b(?:unidades?|unds?|uds?)\b", " un "),
    )
    for pattern, replacement in replacements:
        text = re.sub(pattern, replacement, text)
    text = re.sub(r"(?<=\d),(?=\d)", ".", text)
    text = re.sub(r"[^a-z0-9.]+", " ", text)
    return " ".join(text.split())


MEASUREMENT_PATTERN = re.compile(
    r"(?<![a-z0-9])(?P<value>\d+(?:\.\d+)?)\s*(?P<unit>kg|g|l|ml|cl|un)(?![a-z])"
)
PACK_PATTERN = re.compile(
    r"(?:\bx\s*|\bpack\s+(?:x\s*)?)(?P<value>\d{1,3})"
    r"(?!\s*(?:kg|g|l|ml|cl)\b)(?:\s*un)?\b"
)
MULTIPACK_PATTERN = re.compile(
    r"\b(?P<value>\d{1,2})\s*x\s*\d+(?:\.\d+)?\s*(?:kg|g|l|ml|cl)\b"
)


def parse_measurements(normalized: str) -> list[Measurement]:
    found: dict[tuple[str, int], Measurement] = {}
    for match in MEASUREMENT_PATTERN.finditer(normalized):
        value = float(match.group("value"))
        unit = match.group("unit")
        if value <= 0:
            continue
        if unit == "kg":
            measurement = Measurement("weight_g", value * 1000)
        elif unit == "g":
            measurement = Measurement("weight_g", value)
        elif unit == "l":
            measurement = Measurement("volume_ml", value * 1000)
        elif unit == "cl":
            measurement = Measurement("volume_ml", value * 10)
        elif unit == "ml":
            measurement = Measurement("volume_ml", value)
        else:
            measurement = Measurement("count", value)
        found[(measurement.dimension, round(measurement.value))] = measurement

    for match in PACK_PATTERN.finditer(normalized):
        value = float(match.group("value"))
        if value > 1:
            measurement = Measurement("count", value)
            found[(measurement.dimension, round(measurement.value))] = measurement
    for match in MULTIPACK_PATTERN.finditer(normalized):
        value = float(match.group("value"))
        if value > 1:
            measurement = Measurement("count", value)
            found[(measurement.dimension, round(measurement.value))] = measurement
    return sorted(found.values(), key=lambda item: (item.dimension, item.value))


def tokens_for(normalized: str) -> list[str]:
    return [
        token for token in normalized.split()
        if token not in STOPWORDS
        and not re.fullmatch(r"\d+(?:\.\d+)?", token)
        and token not in {"kg", "g", "l", "ml", "cl", "un"}
    ]


def broad_category(value: str) -> str:
    category = semantic_normalize(value)
    if "bebida" in category:
        return "beverages"
    if "almacen" in category or "kiosko" in category:
        return "pantry"
    if "fresco" in category or "congelado" in category:
        return "fresh"
    if "limpieza" in category or "perfumeria" in category:
        return "cleaning_personal"
    return category


def categories_compatible(left: ProductRecord, right: ProductRecord) -> bool:
    left_category = broad_category(left.source_category)
    right_category = broad_category(right.source_category)
    if left_category == right_category:
        return True
    return {left_category, right_category} <= {"pantry", "fresh", "beverages"}


def measurements_state(left: ProductRecord, right: ProductRecord) -> str:
    left_by_dim = defaultdict(set)
    right_by_dim = defaultdict(set)
    for item in left.measurements:
        left_by_dim[item.dimension].add(round(item.value, 3))
    for item in right.measurements:
        right_by_dim[item.dimension].add(round(item.value, 3))
    common = set(left_by_dim) & set(right_by_dim)
    if "count" not in common:
        one_sided_counts = left_by_dim.get("count", set()) | right_by_dim.get("count", set())
        if any(value > 1 for value in one_sided_counts):
            return "conflict"
    if not common:
        return "unknown"
    for dimension in common:
        left_values = sorted(left_by_dim[dimension])
        right_values = sorted(right_by_dim[dimension])
        compatible = len(left_values) == len(right_values) and all(
            abs(a - b) <= max(0.01, 0.005 * max(a, b))
            for a, b in zip(left_values, right_values)
        )
        if not compatible:
            return "conflict"
    return "exact"


def has_variant_conflict(left: ProductRecord, right: ProductRecord) -> bool:
    left_tokens = set(left.tokens)
    right_tokens = set(right.tokens)
    for family in VARIANT_FAMILIES:
        left_values = left_tokens & family
        right_values = right_tokens & family
        if left_values and right_values and left_values != right_values:
            return True
        if bool(left_values) != bool(right_values) and (left_values | right_values) & ONE_SIDED_CRITICAL_VARIANTS:
            return True
    left_flavors = left_tokens & FLAVORS
    right_flavors = right_tokens & FLAVORS
    if left_flavors and right_flavors and left_flavors != right_flavors:
        return True
    left_wine = left_tokens & WINE_VARIANTS
    right_wine = right_tokens & WINE_VARIANTS
    if left_wine and right_wine and left_wine != right_wine:
        return True
    return bool((left_tokens & BUNDLE_EXTRAS) ^ (right_tokens & BUNDLE_EXTRAS))


def load_workbook_records(path: Path) -> list[ProductRecord]:
    try:
        import openpyxl
    except ImportError as error:  # pragma: no cover - environment guidance
        raise SystemExit("openpyxl is required to read the catalog workbook") from error

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records: list[ProductRecord] = []
    record_id = 0
    for sheet_name in SOURCE_SHEETS:
        worksheet = workbook[sheet_name]
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            category, raw_title, raw_url = row[:3]
            title = repair_text(raw_title)
            url = repair_text(raw_url)
            if not title or not url:
                continue
            source_category = repair_text(category)
            normalized_category = semantic_normalize(source_category)
            if normalized_category not in IN_SCOPE_SOURCE_CATEGORIES[sheet_name]:
                continue
            comparison_title = title_from_url(url) if "�" in title else title
            normalized = semantic_normalize(comparison_title)
            records.append(ProductRecord(
                record_id=record_id,
                store=sheet_name,
                source_category=source_category,
                title=title,
                url=url,
                normalized=normalized,
                tokens=tokens_for(normalized),
                measurements=parse_measurements(normalized),
            ))
            record_id += 1
    return records


def build_idf(records: Iterable[ProductRecord]) -> dict[str, float]:
    records = list(records)
    document_frequency = Counter(token for record in records for token in set(record.tokens))
    count = len(records)
    return {
        token: math.log((count + 1) / (frequency + 1)) + 1
        for token, frequency in document_frequency.items()
    }


def score_pair(left: ProductRecord, right: ProductRecord, idf: dict[str, float]) -> tuple[float, str, str]:
    if not categories_compatible(left, right) or has_variant_conflict(left, right):
        return 0.0, "conflict", "category or variant conflict"
    quantity_state = measurements_state(left, right)
    if quantity_state == "conflict":
        return 0.0, quantity_state, "measurement conflict"

    left_tokens = set(left.tokens)
    right_tokens = set(right.tokens)
    intersection = left_tokens & right_tokens
    if not intersection:
        return 0.0, quantity_state, "no shared product tokens"
    union = left_tokens | right_tokens
    intersection_weight = sum(idf.get(token, 1.0) for token in intersection)
    union_weight = sum(idf.get(token, 1.0) for token in union)
    smaller_weight = min(
        sum(idf.get(token, 1.0) for token in left_tokens),
        sum(idf.get(token, 1.0) for token in right_tokens),
    )
    jaccard = intersection_weight / union_weight if union_weight else 0.0
    containment = intersection_weight / smaller_weight if smaller_weight else 0.0
    character = SequenceMatcher(None, left.normalized, right.normalized).ratio()
    score = 0.42 * containment + 0.33 * jaccard + 0.25 * character
    if quantity_state == "exact":
        score += 0.08
    if left.normalized == right.normalized:
        score += 0.08
    shared_specific = intersection - GENERIC_TOKENS
    if len(shared_specific) >= 2:
        score += 0.03
    score = min(score, 1.0)
    reason = (
        f"weighted tokens={jaccard:.3f}; containment={containment:.3f}; "
        f"text={character:.3f}; quantity={quantity_state}"
    )
    return score, quantity_state, reason


def candidate_edges(records: list[ProductRecord], threshold: float = 0.73) -> list[MatchEdge]:
    idf = build_idf(records)
    records_by_store = {store: [record for record in records if record.store == store] for store in SOURCE_SHEETS}
    indexes: dict[str, dict[str, set[int]]] = {}
    by_id = {record.record_id: record for record in records}
    for store, store_records in records_by_store.items():
        token_index: dict[str, set[int]] = defaultdict(set)
        for record in store_records:
            for token in set(record.tokens):
                token_index[token].add(record.record_id)
        indexes[store] = token_index

    directed_best: dict[tuple[int, str], tuple[float, float, MatchEdge]] = {}
    store_pairs = [
        (SOURCE_SHEETS[left], SOURCE_SHEETS[right])
        for left in range(len(SOURCE_SHEETS))
        for right in range(left + 1, len(SOURCE_SHEETS))
    ]
    for left_store, right_store in store_pairs:
        print(f"matching {left_store} -> {right_store}", file=sys.stderr, flush=True)
        right_index = indexes[right_store]
        for left in records_by_store[left_store]:
            specific_tokens = set(left.tokens) - GENERIC_TOKENS
            ranked_tokens = sorted(specific_tokens, key=lambda token: idf.get(token, 0), reverse=True)
            selective_tokens = [
                token for token in ranked_tokens
                if len(right_index.get(token, ())) <= 250
            ][:3]
            if not selective_tokens:
                selective_tokens = sorted(
                    set(left.tokens), key=lambda token: idf.get(token, 0), reverse=True
                )[:1]
            candidate_ids: set[int] = set()
            for token in selective_tokens:
                candidate_ids.update(right_index.get(token, ()))
            scored: list[MatchEdge] = []
            for candidate_id in candidate_ids:
                right = by_id[candidate_id]
                if not categories_compatible(left, right):
                    continue
                shared_specific = (set(left.tokens) & set(right.tokens)) - GENERIC_TOKENS
                if len(shared_specific) < 2 and measurements_state(left, right) != "exact":
                    continue
                score, quantity_state, reason = score_pair(left, right, idf)
                if score >= threshold:
                    scored.append(MatchEdge(left.record_id, right.record_id, score, quantity_state, reason))
            scored.sort(key=lambda edge: edge.score, reverse=True)
            if scored:
                second = scored[1].score if len(scored) > 1 else 0.0
                directed_best[(left.record_id, right_store)] = (scored[0].score, second, scored[0])

    reciprocal: dict[tuple[int, int], MatchEdge] = {}
    for (left_id, right_store), (score, second, edge) in directed_best.items():
        right_id = edge.right_id
        reverse = directed_best.get((right_id, by_id[left_id].store))
        is_reciprocal = reverse is not None and reverse[2].right_id == left_id
        decisive = score >= 0.93 and score - second >= 0.06
        if is_reciprocal or decisive:
            key = tuple(sorted((left_id, right_id)))
            reciprocal[key] = edge
    return list(reciprocal.values())


def connected_groups(records: list[ProductRecord], edges: list[MatchEdge]) -> list[dict]:
    by_id = {record.record_id: record for record in records}
    parent = {record.record_id: record.record_id for record in records}

    def find(value: int) -> int:
        while parent[value] != value:
            parent[value] = parent[parent[value]]
            value = parent[value]
        return value

    def union(left: int, right: int) -> None:
        left_root, right_root = find(left), find(right)
        if left_root != right_root:
            parent[right_root] = left_root

    for edge in edges:
        union(edge.left_id, edge.right_id)

    component_ids: dict[int, set[int]] = defaultdict(set)
    for edge in edges:
        component_ids[find(edge.left_id)].update((edge.left_id, edge.right_id))
    edges_by_root: dict[int, list[MatchEdge]] = defaultdict(list)
    for edge in edges:
        edges_by_root[find(edge.left_id)].append(edge)

    output: list[dict] = []
    for root, ids in component_ids.items():
        component = [by_id[record_id] for record_id in ids]
        store_counts = Counter(record.store for record in component)
        duplicate_store = any(count > 1 for count in store_counts.values())
        active_count = sum(record.store in ACTIVE_STORES for record in component)
        component_edges = edges_by_root[root]
        min_score = min(edge.score for edge in component_edges)
        average_score = sum(edge.score for edge in component_edges) / len(component_edges)
        quantities_known = all(record.measurements for record in component)
        category = classify_cuanto_category(component)
        component_idf = build_idf(component)
        pairwise_compatible = all(
            score_pair(left, right, component_idf)[0] > 0
            for index, left in enumerate(component)
            for right in component[index + 1:]
            if left.store != right.store
        )
        confidence = "review"
        if (
            not duplicate_store
            and active_count >= 2
            and category
            and quantities_known
            and pairwise_compatible
            and min_score >= 0.84
            and average_score >= 0.87
        ):
            confidence = "high"
        elif not duplicate_store and active_count >= 2 and category and min_score >= 0.76:
            confidence = "medium"
        output.append({
            "confidence": confidence,
            "cuanto_category": category,
            "active_store_count": active_count,
            "duplicate_store": duplicate_store,
            "minimum_score": round(min_score, 4),
            "average_score": round(average_score, 4),
            "records": [record_to_output(record) for record in sorted(component, key=lambda item: item.store)],
            "edges": [asdict(edge) for edge in component_edges],
        })
    output.sort(key=lambda group: (-group["active_store_count"], -group["average_score"], group["records"][0]["title"]))
    return output


def classify_cuanto_category(records: list[ProductRecord]) -> str | None:
    categories = {broad_category(record.source_category) for record in records}
    tokens = set(token for record in records for token in record.tokens)
    if "beverages" in categories:
        return "Bebidas"
    if "cleaning_personal" in categories:
        return "Limpieza" if tokens & CLEANING_TERMS else None
    if categories == {"fresh"} and tokens & PREPARED_FOOD_TERMS:
        return None
    if "fresh" in categories and tokens & DAIRY_TERMS:
        return "Lácteos"
    if "fresh" in categories and tokens & MEAT_TERMS:
        return "Carnes"
    if categories <= {"pantry", "fresh"} and "pantry" in categories:
        return "Almacén"
    return None


def record_to_output(record: ProductRecord) -> dict:
    return {
        "record_id": record.record_id,
        "store": record.store,
        "source_category": record.source_category,
        "title": record.title,
        "url": record.url,
        "normalized": record.normalized,
        "measurements": [asdict(item) for item in record.measurements],
    }


def summarize(groups: list[dict], records: list[ProductRecord]) -> dict:
    return {
        "source_records": len(records),
        "groups": len(groups),
        "confidence": dict(Counter(group["confidence"] for group in groups)),
        "categories": dict(Counter(group["cuanto_category"] or "excluded" for group in groups)),
        "active_store_counts": dict(Counter(str(group["active_store_count"]) for group in groups)),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--summary", type=Path)
    args = parser.parse_args()

    records = load_workbook_records(args.workbook)
    edges = candidate_edges(records)
    groups = connected_groups(records, edges)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(groups, ensure_ascii=False, indent=2), encoding="utf-8")
    summary = summarize(groups, records)
    if args.summary:
        args.summary.parent.mkdir(parents=True, exist_ok=True)
        args.summary.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
